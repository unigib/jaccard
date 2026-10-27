import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
)

from app import (
    compare_to_corpus,
    extract_document_text,
    jaccard_similarity,
    make_ngrams,
    normalize_text,
)


PROJECT_DIR = Path(__file__).resolve().parent


class JaccardAlgorithmTests(unittest.TestCase):
    def test_normalization_lowercases_and_removes_punctuation(self):
        self.assertEqual(
            normalize_text("Hello, WORLD! It's fine_here."),
            "hello world it s fine here",
        )

    def test_make_ngrams_returns_sliding_unique_trigrams(self):
        self.assertEqual(
            make_ngrams("The quick brown fox jumps"),
            {
                "the quick brown",
                "quick brown fox",
                "brown fox jumps",
            },
        )

    def test_jaccard_similarity_uses_intersection_over_union(self):
        self.assertAlmostEqual(
            jaccard_similarity({"a", "b", "c"}, {"b", "c", "d"}),
            50.0,
        )

    def test_two_empty_sets_have_zero_similarity(self):
        self.assertEqual(jaccard_similarity(set(), set()), 0.0)

    def test_comparison_reports_progress_for_each_reference_file(self):
        with TemporaryDirectory() as temp_dir:
            corpus = Path(temp_dir)
            (corpus / "second.txt").write_text(
                "A second example reference document.", encoding="utf-8"
            )
            (corpus / "first.txt").write_text(
                "A first example reference document.", encoding="utf-8"
            )
            progress = []

            comparisons = compare_to_corpus(
                "A first example reference document.",
                corpus,
                on_progress=lambda position, total, filename: progress.append(
                    (position, total, filename)
                ),
            )

        self.assertEqual(
            progress,
            [(1, 2, "first.txt"), (2, 2, "second.txt")],
        )
        self.assertEqual(len(comparisons), 2)

    def test_demo_inputs_cover_original_and_suspect_results(self):
        original_text = (PROJECT_DIR / "demo_inputs" / "original.txt").read_text(
            encoding="utf-8"
        )
        suspect_text = (PROJECT_DIR / "demo_inputs" / "suspect.txt").read_text(
            encoding="utf-8"
        )

        original = compare_to_corpus(original_text)
        suspect = compare_to_corpus(suspect_text)

        self.assertGreaterEqual(len(original), 5)
        self.assertLess(original[0].score, 30)
        self.assertEqual(suspect[0].filename, "article_1_climate.txt")
        self.assertGreaterEqual(suspect[0].score, 30)

    def test_extracts_text_from_txt_file(self):
        text_path = PROJECT_DIR / "demo_inputs" / "original.txt"
        self.assertEqual(
            extract_document_text(text_path),
            text_path.read_text(encoding="utf-8"),
        )

    def test_rejects_unsupported_document_type(self):
        with self.assertRaisesRegex(ValueError, "Choose a .txt, .docx, or .pdf"):
            extract_document_text(PROJECT_DIR / "unsupported.rtf")

    def test_extracts_docx_paragraphs_and_table_content(self):
        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "document.docx"
            document = Document()
            document.add_paragraph("DOCX extraction test text.")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "first cell"
            table.cell(0, 1).text = "second cell"
            document.save(path)

            extracted = extract_document_text(path)

        self.assertIn("DOCX extraction test text.", extracted)
        self.assertIn("first cell\tsecond cell", extracted)

    def test_extracts_text_from_pdf(self):
        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "document.pdf"
            writer = PdfWriter()
            page = writer.add_blank_page(width=612, height=792)
            font = writer._add_object(
                DictionaryObject(
                    {
                        NameObject("/Type"): NameObject("/Font"),
                        NameObject("/Subtype"): NameObject("/Type1"),
                        NameObject("/BaseFont"): NameObject("/Helvetica"),
                    }
                )
            )
            page[NameObject("/Resources")] = DictionaryObject(
                {
                    NameObject("/Font"): DictionaryObject(
                        {NameObject("/F1"): font}
                    )
                }
            )
            content = DecodedStreamObject()
            content.set_data(b"BT /F1 12 Tf 50 700 Td (PDF extraction test.) Tj ET")
            page[NameObject("/Contents")] = writer._add_object(content)
            with path.open("wb") as pdf_file:
                writer.write(pdf_file)

            extracted = extract_document_text(path)

        self.assertEqual(extracted.strip(), "PDF extraction test.")

    def test_pdf_without_selectable_text_explains_ocr_limitation(self):
        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "image-only.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=612, height=792)
            with path.open("wb") as pdf_file:
                writer.write(pdf_file)

            with self.assertRaisesRegex(ValueError, "OCR is not included"):
                extract_document_text(path)


if __name__ == "__main__":
    unittest.main()
