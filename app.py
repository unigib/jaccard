"""Tkinter demonstration of trigram-based Jaccard document similarity."""

from __future__ import annotations

import re
import queue
import threading
import tkinter as tk
import zipfile
from dataclasses import dataclass
from pathlib import Path
from tkinter import filedialog, messagebox, ttk
from typing import Callable


APP_DIR = Path(__file__).resolve().parent
CORPUS_DIR = APP_DIR / "corpus"
LOGO_PATH = APP_DIR / "logo.png"
TRIGRAM_SIZE = 3
DEFAULT_THRESHOLD = 30.0


@dataclass(frozen=True)
class Comparison:
    filename: str
    score: float
    shared_trigrams: tuple[str, ...]


def normalize_text(text: str) -> str:
    """Lowercase text and replace punctuation with spaces."""
    cleaned = re.sub(r"[^\w\s]|_", " ", text.lower(), flags=re.UNICODE)
    return " ".join(cleaned.split())


def _ngrams_from_normalized(normalized_text: str, n: int) -> set[str]:
    words = normalized_text.split()
    return {
        " ".join(words[index : index + n])
        for index in range(len(words) - n + 1)
    }


def make_ngrams(text: str, n: int = TRIGRAM_SIZE) -> set[str]:
    """Return unique sliding word n-grams from normalized or raw text."""
    if n < 1:
        raise ValueError("N-gram size must be at least 1.")
    return _ngrams_from_normalized(normalize_text(text), n)


def jaccard_similarity(first: set[str], second: set[str]) -> float:
    """Return the Jaccard similarity of two sets as a percentage."""
    union = first | second
    if not union:
        return 0.0
    return len(first & second) / len(union) * 100


def compare_to_corpus(
    text: str,
    corpus_dir: Path = CORPUS_DIR,
    on_progress: Callable[[int, int, str], None] | None = None,
    on_phase: Callable[[str, int, int, str], None] | None = None,
    trigram_cache: dict[Path, tuple[int, int, set[str]]] | None = None,
) -> list[Comparison]:
    """Compare text against every UTF-8 .txt file in the corpus."""
    if on_phase is not None:
        on_phase("Phase 1/4 — Normalizing submitted document", 0, 1, "")
    normalized_text = normalize_text(text)
    if on_phase is not None:
        on_phase("Phase 2/4 — Creating submitted trigrams", 0, 1, "")
    query_trigrams = _ngrams_from_normalized(normalized_text, TRIGRAM_SIZE)
    if not query_trigrams:
        raise ValueError("Enter at least three words to create a trigram.")

    documents = sorted(corpus_dir.glob("*.txt"), key=lambda path: path.name.casefold())
    if not documents:
        raise FileNotFoundError(
            f"No .txt reference documents were found in {corpus_dir}."
        )

    cache = trigram_cache if trigram_cache is not None else {}
    document_trigrams_by_path: dict[Path, set[str]] = {}
    cache_hits = 0
    for path in documents:
        stat = path.stat()
        cached = cache.get(path)
        if cached is not None and cached[:2] == (stat.st_mtime_ns, stat.st_size):
            document_trigrams_by_path[path] = cached[2]
            cache_hits += 1

    if on_phase is not None:
        on_phase(
            "Phase 3/4 — Preparing reference trigrams",
            0,
            len(documents),
            f"{cache_hits} of {len(documents)} references already indexed",
        )
    for position, path in enumerate(documents, start=1):
        if path not in document_trigrams_by_path:
            stat = path.stat()
            document_trigrams = _ngrams_from_normalized(
                normalize_text(path.read_text(encoding="utf-8")),
                TRIGRAM_SIZE,
            )
            cache[path] = (stat.st_mtime_ns, stat.st_size, document_trigrams)
            document_trigrams_by_path[path] = document_trigrams
            detail = f"Indexing {position} of {len(documents)}: {path.name}"
        else:
            detail = f"Using cached trigrams {position} of {len(documents)}: {path.name}"
        if on_phase is not None:
            on_phase(
                "Phase 3/4 — Preparing reference trigrams",
                position,
                len(documents),
                detail,
            )

    current_paths = set(documents)
    for cached_path in tuple(cache):
        if cached_path not in current_paths:
            del cache[cached_path]

    if on_phase is not None:
        on_phase(
            "Phase 4/4 — Comparing Jaccard similarity",
            0,
            len(documents),
            "",
        )
    comparisons: list[Comparison] = []
    for position, path in enumerate(documents, start=1):
        document_trigrams = document_trigrams_by_path[path]
        shared_set = query_trigrams & document_trigrams
        union_size = len(query_trigrams) + len(document_trigrams) - len(shared_set)
        score = len(shared_set) / union_size * 100 if union_size else 0.0
        shared = tuple(sorted(shared_set))
        comparisons.append(
            Comparison(
                filename=path.name,
                score=score,
                shared_trigrams=shared,
            )
        )
        if on_progress is not None:
            on_progress(position, len(documents), path.name)
        if on_phase is not None:
            on_phase(
                "Phase 4/4 — Comparing Jaccard similarity",
                position,
                len(documents),
                f"Comparing {position} of {len(documents)}: {path.name}",
            )
    return sorted(comparisons, key=lambda result: (-result.score, result.filename.casefold()))


def extract_document_text(path: Path) -> str:
    """Extract text from a .txt, .docx, or text-based .pdf document."""
    suffix = path.suffix.casefold()
    if suffix == ".txt":
        text = path.read_text(encoding="utf-8")
    elif suffix == ".docx":
        try:
            from docx import Document
        except ImportError as error:
            raise RuntimeError(
                "DOCX support requires python-docx. Install dependencies with "
                "'python -m pip install -r requirements.txt'."
            ) from error

        try:
            document = Document(path)
        except (OSError, ValueError, zipfile.BadZipFile) as error:
            raise ValueError(f"Could not read the DOCX document: {error}") from error
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        for table in document.tables:
            for row in table.rows:
                parts.append("\t".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    elif suffix == ".pdf":
        try:
            from pypdf import PdfReader
            from pypdf.errors import FileNotDecryptedError, PdfReadError
        except ImportError as error:
            raise RuntimeError(
                "PDF support requires pypdf. Install dependencies with "
                "'python -m pip install -r requirements.txt'."
            ) from error

        try:
            reader = PdfReader(path)
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        except (FileNotDecryptedError, OSError, PdfReadError) as error:
            raise ValueError(f"Could not extract text from the PDF: {error}") from error
    else:
        raise ValueError("Choose a .txt, .docx, or .pdf file.")

    if not text.strip():
        if suffix == ".pdf":
            raise ValueError(
                "No selectable text was found in this PDF. It may be a scanned "
                "document; OCR is not included."
            )
        raise ValueError("The selected document does not contain any text.")
    return text


class JaccardApp(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title("University of Gibraltar | Document Similarity Checker")
        self.geometry("1000x760")
        self.minsize(800, 620)
        self.configure(bg="#f1f5f9")

        self.threshold = tk.StringVar(value=str(int(DEFAULT_THRESHOLD)))
        self._events: queue.Queue[tuple] = queue.Queue()
        self._check_worker: threading.Thread | None = None
        self._trigram_cache: dict[Path, tuple[int, int, set[str]]] = {}
        self._build_ui()
        self.after(100, self._process_check_events)

    def _build_ui(self) -> None:
        style = ttk.Style(self)
        style.configure("Treeview", rowheight=27, font=("Segoe UI", 10))
        style.configure("Treeview.Heading", font=("Segoe UI", 10, "bold"))
        style.configure("Accent.TButton", font=("Segoe UI", 10, "bold"))

        header = tk.Frame(
            self, bg="#ffffff", padx=24, pady=15, highlightbackground="#e2e8f0",
            highlightthickness=1,
        )
        header.pack(fill="x")
        self.logo_image = tk.PhotoImage(master=self, file=str(LOGO_PATH)).subsample(2, 2)
        tk.Label(header, image=self.logo_image, bg="#ffffff").pack(
            side="left", padx=(0, 20)
        )
        brand_text = tk.Frame(header, bg="#ffffff")
        brand_text.pack(side="left", fill="x", expand=True)
        tk.Label(
            brand_text,
            text="UNIVERSITY OF GIBRALTAR",
            font=("Segoe UI", 11, "bold"),
            fg="#a6192e",
            bg="#ffffff",
        ).pack(anchor="w")
        tk.Label(
            brand_text,
            text="Document Similarity Checker",
            font=("Segoe UI", 19, "bold"),
            fg="#1e293b",
            bg="#ffffff",
        ).pack(anchor="w", pady=(3, 0))
        tk.Label(
            brand_text,
            text="Normalize text  →  create word trigrams  →  compare with Jaccard similarity",
            font=("Segoe UI", 10),
            fg="#475569",
            bg="#ffffff",
        ).pack(anchor="w", pady=(4, 0))

        body = ttk.Frame(self, padding=(22, 16))
        body.pack(fill="both", expand=True)
        ttk.Label(body, text="Document to check", font=("Segoe UI", 12, "bold")).pack(
            anchor="w"
        )
        ttk.Label(
            body,
            text="Paste text or load a .txt, .docx, or text-based .pdf file. Extracted text appears here and can be saved as .txt.",
            wraplength=900,
        ).pack(anchor="w", pady=(3, 8))

        input_frame = ttk.Frame(body)
        input_frame.pack(fill="both", expand=True)
        self.input_text = tk.Text(
            input_frame,
            height=9,
            wrap="word",
            font=("Segoe UI", 10),
            undo=True,
            padx=10,
            pady=9,
            borderwidth=1,
            relief="solid",
        )
        input_scroll = ttk.Scrollbar(
            input_frame, orient="vertical", command=self.input_text.yview
        )
        self.input_text.configure(yscrollcommand=input_scroll.set)
        self.input_text.pack(side="left", fill="both", expand=True)
        input_scroll.pack(side="right", fill="y")

        controls = ttk.Frame(body, padding=(0, 10, 0, 12))
        controls.pack(fill="x")
        self.load_button = ttk.Button(
            controls, text="Load document", command=self._load_file
        )
        self.load_button.pack(side="left")
        self.save_button = ttk.Button(
            controls, text="Save text as .txt", command=self._save_text
        )
        self.save_button.pack(side="left", padx=(8, 0))
        self.check_button = ttk.Button(
            controls,
            text="Check document",
            command=self._check_document,
            style="Accent.TButton",
        )
        self.check_button.pack(side="left", padx=(8, 0))
        self.clear_button = ttk.Button(
            controls, text="Clear", command=self._clear
        )
        self.clear_button.pack(side="left", padx=(8, 0))
        ttk.Label(controls, text="Flag threshold:").pack(side="left", padx=(20, 5))
        self.threshold_entry = ttk.Entry(
            controls, textvariable=self.threshold, width=6
        )
        self.threshold_entry.pack(side="left")
        ttk.Label(controls, text="%").pack(side="left", padx=(3, 0))

        self.status = tk.Label(
            body,
            text="Ready — enter at least three words and run a check.",
            anchor="w",
            font=("Segoe UI", 11, "bold"),
            padx=12,
            pady=9,
            bg="#e2e8f0",
            fg="#334155",
        )
        self.status.pack(fill="x", pady=(0, 10))
        self.progress = ttk.Progressbar(body, mode="determinate")
        self.progress.pack(fill="x", pady=(0, 10))
        self.progress.pack_forget()

        ttk.Label(body, text="Reference document scores", font=("Segoe UI", 11, "bold")).pack(
            anchor="w", pady=(0, 5)
        )
        result_frame = ttk.Frame(body)
        result_frame.pack(fill="both", expand=True)
        self.results = ttk.Treeview(
            result_frame,
            columns=("document", "score", "shared"),
            show="headings",
            height=5,
        )
        self.results.heading("document", text="Reference document")
        self.results.heading("score", text="Similarity")
        self.results.heading("shared", text="Shared trigrams")
        self.results.column("document", width=220, anchor="w")
        self.results.column("score", width=110, anchor="center")
        self.results.column("shared", width=500, anchor="w")
        result_scroll = ttk.Scrollbar(
            result_frame, orient="vertical", command=self.results.yview
        )
        self.results.configure(yscrollcommand=result_scroll.set)
        self.results.pack(side="left", fill="both", expand=True)
        result_scroll.pack(side="right", fill="y")
        self.results.bind("<<TreeviewSelect>>", self._show_selected_phrases)

        self.phrase_label = ttk.Label(
            body,
            text="Select a result row to see its complete list of shared trigrams.",
            wraplength=920,
            justify="left",
        )
        self.phrase_label.pack(anchor="w", fill="x", pady=(8, 0))
        self._comparisons: dict[str, Comparison] = {}

    def _load_file(self) -> None:
        filename = filedialog.askopenfilename(
            title="Load document",
            filetypes=(
                ("Supported documents", "*.txt *.docx *.pdf"),
                ("Text files", "*.txt"),
                ("Word documents", "*.docx"),
                ("PDF documents", "*.pdf"),
                ("All files", "*.*"),
            ),
        )
        if not filename:
            return
        try:
            content = extract_document_text(Path(filename))
        except (OSError, UnicodeError, ValueError, RuntimeError) as error:
            messagebox.showerror("Could not load file", str(error), parent=self)
            return
        self.input_text.delete("1.0", "end")
        self.input_text.insert("1.0", content)

    def _save_text(self) -> None:
        content = self.input_text.get("1.0", "end-1c")
        if not content.strip():
            messagebox.showwarning(
                "No text to save",
                "Load or enter a document before saving it as .txt.",
                parent=self,
            )
            return
        filename = filedialog.asksaveasfilename(
            title="Save extracted text",
            defaultextension=".txt",
            filetypes=(("Text files", "*.txt"),),
        )
        if not filename:
            return
        try:
            Path(filename).write_text(content, encoding="utf-8")
        except OSError as error:
            messagebox.showerror("Could not save text", str(error), parent=self)
            return
        messagebox.showinfo(
            "Text saved", f"Extracted text saved to:\n{filename}", parent=self
        )

    def _check_document(self) -> None:
        if self._check_worker is not None and self._check_worker.is_alive():
            return
        try:
            threshold = float(self.threshold.get())
            if not 0 <= threshold <= 100:
                raise ValueError("Threshold must be between 0 and 100.")
        except ValueError as error:
            messagebox.showwarning("Check the input", str(error), parent=self)
            return

        text = self.input_text.get("1.0", "end-1c")
        self._set_check_controls_enabled(False)
        self.progress.configure(mode="determinate", maximum=1, value=0)
        self.progress.pack(fill="x", pady=(0, 10))
        self.status.configure(
            text="Starting document check…",
            bg="#dbeafe",
            fg="#1e40af",
        )
        self._check_worker = threading.Thread(
            target=self._compare_in_background,
            args=(text, threshold),
            daemon=True,
        )
        self._check_worker.start()

    def _compare_in_background(self, text: str, threshold: float) -> None:
        try:
            comparisons = compare_to_corpus(
                text,
                on_progress=lambda position, total, filename: self._events.put(
                    ("progress", position, total, filename)
                ),
                on_phase=lambda phase, position, total, detail: self._events.put(
                    ("phase", phase, position, total, detail)
                ),
                trigram_cache=self._trigram_cache,
            )
        except (OSError, UnicodeError, ValueError) as error:
            self._events.put(("error", str(error)))
            return
        self._events.put(("complete", comparisons, threshold))

    def _process_check_events(self) -> None:
        while True:
            try:
                event = self._events.get_nowait()
            except queue.Empty:
                break

            if event[0] == "progress":
                _, position, total, filename = event
                self.progress.configure(
                    mode="determinate", maximum=total, value=position
                )
                self.status.configure(
                    text=f"Phase 4/4 — Comparing {position} of {total}: {filename}",
                    bg="#dbeafe",
                    fg="#1e40af",
                )
            elif event[0] == "phase":
                _, phase, position, total, detail = event
                self.status.configure(
                    text=f"{phase}{': ' + detail if detail else ''}",
                    bg="#dbeafe",
                    fg="#1e40af",
                )
                self.progress.configure(
                    mode="determinate",
                    maximum=max(total, 1),
                    value=position,
                )
            elif event[0] == "error":
                self._finish_check()
                messagebox.showerror(
                    "Could not scan corpus", event[1], parent=self
                )
                self.status.configure(
                    text="Check could not be completed.",
                    bg="#fee2e2",
                    fg="#991b1b",
                )
            elif event[0] == "complete":
                _, comparisons, threshold = event
                self._finish_check()
                self._display_comparisons(comparisons, threshold)

        self.after(100, self._process_check_events)

    def _finish_check(self) -> None:
        self.progress.stop()
        self.progress.pack_forget()
        self._set_check_controls_enabled(True)
        self._check_worker = None

    def _set_check_controls_enabled(self, enabled: bool) -> None:
        state = "normal" if enabled else "disabled"
        for control in (
            self.load_button,
            self.save_button,
            self.check_button,
            self.clear_button,
            self.threshold_entry,
        ):
            control.configure(state=state)

    def _display_comparisons(
        self, comparisons: list[Comparison], threshold: float
    ) -> None:
        for item in self.results.get_children():
            self.results.delete(item)
        self._comparisons.clear()
        for comparison in comparisons:
            item_id = self.results.insert(
                "",
                "end",
                values=(
                    comparison.filename,
                    f"{comparison.score:.1f}%",
                    ", ".join(comparison.shared_trigrams[:3])
                    + (" …" if len(comparison.shared_trigrams) > 3 else ""),
                ),
            )
            self._comparisons[item_id] = comparison

        best = comparisons[0]
        if best.score >= threshold:
            self.status.configure(
                text=f"SUSPECT PLAGIARISM — {best.score:.1f}% match with {best.filename}",
                bg="#fee2e2",
                fg="#991b1b",
            )
        else:
            self.status.configure(
                text=f"ORIGINAL — highest match is {best.score:.1f}% with {best.filename}",
                bg="#dcfce7",
                fg="#166534",
            )
        self.results.selection_set(self.results.get_children()[0])
        self._show_selected_phrases()

    def _show_selected_phrases(self, _event: tk.Event | None = None) -> None:
        selection = self.results.selection()
        if not selection:
            return
        comparison = self._comparisons.get(selection[0])
        if comparison is None:
            return
        if comparison.shared_trigrams:
            phrases = "; ".join(comparison.shared_trigrams)
            self.phrase_label.configure(
                text=f"All shared trigrams with {comparison.filename}: {phrases}"
            )
        else:
            self.phrase_label.configure(
                text=f"No shared trigrams with {comparison.filename}."
            )

    def _clear(self) -> None:
        self.input_text.delete("1.0", "end")
        for item in self.results.get_children():
            self.results.delete(item)
        self._comparisons.clear()
        self.status.configure(
            text="Ready — enter at least three words and run a check.",
            bg="#e2e8f0",
            fg="#334155",
        )
        self.phrase_label.configure(
            text="Select a result row to see its complete list of shared trigrams."
        )


if __name__ == "__main__":
    JaccardApp().mainloop()
